"""PLAYGROUND_CATALOG(v7)을 원본 시트용 XLSX로 바꾼다. Drive에 올릴 때 Google 시트로 변환된다.

사용: python tools/make-sheet.py <playground-v7.js> <out.xlsx> [gw-track.json]
"""
import json
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

TUT_COLS = ["id", "publish_state", "category", "code", "title", "description", "tags", "difficulty",
            "difficultyBasis", "minutes", "timeBasis", "mode", "engine", "practiceUrl", "sourceUrl",
            "scopeNote", "expected", "recovery", "hints", "conditional", "status", "sourceKind", "courseId",
            "sourceStatus", "volumeIds", "simulation", "rights_json", "blockers_json", "resources_json",
            "cover_media_id", "updated_by", "updated_at"]
STEP_COLS = ["tutorial_id", "order", "code", "title", "instruction", "expected", "recovery", "kind", "sourceType",
             "assessmentBasis", "sourceUrl", "practiceUrl", "media_id", "media_caption", "extra_json"]
MEDIA_COLS = ["media_id", "drive_file_id", "type", "display", "alt", "caption", "source", "rights", "public_url",
              "uploaded_by", "uploaded_at", "used_in"]
LINK_COLS = ["link_id", "tutorial_id", "url", "audience", "channel", "created_by", "created_at", "note"]
LEARNER_COLS = ["email", "name", "group", "source", "expires", "note", "added_at"]
ACCESS_COLS = ["ts", "email", "trace", "action", "paths", "tutorial_id", "result", "ms"]
LOG_COLS = ["ts", "user", "action", "target", "detail"]
VERSION_COLS = ["version", "file_id", "published_by", "published_at", "items", "note"]

KNOWN_STEP = {"title", "instruction", "expected", "recovery", "kind", "sourceType", "assessmentBasis", "code",
              "sourceUrl", "practiceUrl", "media"}


def j(v):
    return "" if v in (None, [], {}) else json.dumps(v, ensure_ascii=False)


def b(v):
    return "TRUE" if v else "FALSE"


def tut_row(it, state):
    return [it["id"], state, it.get("category"), it.get("code"), it.get("title"), it.get("description"),
            ", ".join(it.get("tags") or []), it.get("difficulty"), it.get("difficultyBasis"), it.get("minutes"),
            it.get("timeBasis"), it.get("mode"), it.get("engine") or "", it.get("practiceUrl") or "",
            it.get("sourceUrl") or "", it.get("scopeNote") or "", it.get("expected") or "", it.get("recovery") or "",
            "\n".join(it.get("hints") or []), b(it.get("conditional")), it.get("status") or "",
            it.get("sourceKind") or "", it.get("courseId") or "", it.get("sourceStatus") or "",
            ", ".join(it.get("volumeIds") or []), b(it.get("simulation")), j(it.get("rights")), j(it.get("blockers")),
            j(it.get("resources")), "", "import", datetime.now().strftime("%Y-%m-%d %H:%M")]


def step_rows(it):
    rows = []
    for i, s in enumerate(it.get("steps") or [], 1):
        extra = {k: v for k, v in s.items() if k not in KNOWN_STEP}
        rows.append([it["id"], i, s.get("code"), s.get("title"), s.get("instruction"), s.get("expected") or "",
                     s.get("recovery") or "", s.get("kind") or "manual", s.get("sourceType") or "",
                     s.get("assessmentBasis") or "", s.get("sourceUrl") or "", s.get("practiceUrl") or "", "", "",
                     j(extra)])
    return rows


def sheet(wb, name, cols, rows, widths=None):
    ws = wb.create_sheet(name)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3B4A6B")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "B2" if name in ("tutorials", "steps") else "A2"
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    return ws


def main():
    src, out = sys.argv[1], sys.argv[2]
    text = open(src, encoding="utf-8").read()
    cat = json.loads(text[text.index("=") + 1:].strip().rstrip(";"))
    gw = json.load(open(sys.argv[3], encoding="utf-8")) if len(sys.argv) > 3 else []
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "안내"
    for line in [
        ["한 걸음 서재 원본 시트"],
        ["이 시트가 사이트 내용의 원본입니다. 직접 고쳐도 되지만, 관리자 웹앱을 쓰면 검사와 미리보기가 함께 됩니다."],
        ["publish_state: 초안 / 검수중 / 게시 / 보관. '게시'인 행만 사이트에 나갑니다(운영자가 게시 버튼을 눌러야 반영)."],
        ["tags·volumeIds는 쉼표로, hints는 줄바꿈으로 구분합니다. *_json 열은 형식을 지켜야 합니다."],
        ["steps 탭: tutorial_id + order 순서대로 보입니다. media_id는 media 탭의 ID입니다."],
        ["권한은 Drive 공유로 정합니다. 이 시트 편집자 = 편집자, 댓글 작성자 = 검수자, 40_게시본 폴더 편집자 = 운영자."],
    ]:
        ws0.append(line)
    ws0.column_dimensions["A"].width = 120
    items = cat["items"]
    sheet(wb, "tutorials", TUT_COLS,
          [tut_row(it, "게시") for it in items] + [tut_row(it, "초안") for it in gw],
          [26, 10, 12, 8, 40, 50, 30])
    srows = [r for it in items for r in step_rows(it)] + [r for it in gw for r in step_rows(it)]
    sheet(wb, "steps", STEP_COLS, srows, [26, 6, 10, 12, 60, 40, 40])
    sheet(wb, "media", MEDIA_COLS, [], [14, 30, 8, 12, 40])
    sheet(wb, "links", LINK_COLS, [])
    sheet(wb, "learners", LEARNER_COLS, [["soundfury37@gmail.com", "운영자", "운영", "manual", "", "운영자 계정", datetime.now().strftime("%Y-%m-%d")]], [32, 14, 10, 14, 12])
    sheet(wb, "access_log", ACCESS_COLS, [], [18, 28, 10, 10, 40, 24, 8, 6])
    sheet(wb, "log", LOG_COLS, [])
    sheet(wb, "versions", VERSION_COLS, [[1, "", "import", datetime.now().strftime("%Y-%m-%d %H:%M"), len(items), "기존 v7 사이트에서 가져온 첫 게시본"]])
    settings = [
        ["categories_json", j(cat.get("categories") + ["구글 워크스페이스"])],
        ["difficultyOptions_json", j(cat.get("difficultyOptions"))],
        ["policies_json", j(cat.get("policies"))],
        ["volumes_json", j(cat.get("volumes"))],
        ["epub_json", j(cat.get("epub"))],
        ["sourceHashes_json", j(cat.get("sourceHashes"))],
        ["stats_extra_json", j({k: v for k, v in cat["stats"].items() if k not in ("items", "practice", "guide")})],
        ["preview_steps", 2],
        ["quota_per_6h", 120],
        ["gemini_model", "gemini-2.5-flash"],
    ]
    ws = sheet(wb, "settings", ["key", "value"], settings, [26, 100])
    # 드롭다운
    tut = wb["tutorials"]
    dv = DataValidation(type="list", formula1='"초안,검수중,게시,보관"', allow_blank=False)
    tut.add_data_validation(dv)
    dv.add(f"B2:B{len(items) + len(gw) + 500}")
    wb.save(out)
    print(json.dumps({"tutorials": len(items) + len(gw), "steps": len(srows), "gw": len(gw)}))


if __name__ == "__main__":
    main()
